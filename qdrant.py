import os
import io
from datetime import datetime
import asyncio
from nicegui import ui, events
from qdrant_client import QdrantClient, models
from qdrant_client.models import PointStruct, Filter
from openai import OpenAI
import PyPDF2
import docx
from langchain.text_splitter import RecursiveCharacterTextSplitter
import uuid
from dotenv import load_dotenv
import json

# Assuming styles.py exists; if not, replace with inline styles
from styles import apply_custom_styles

# Load environment variables from .env file
load_dotenv()

# Configuration
QDRANT_HOST     = os.getenv("QDRANT_HOST", "http://localhost:6333")
QDRANT_API_KEY  = os.getenv("QDRANT_API_KEY")
OPENAI_API_KEY  = os.getenv("OPENAI_API_KEY")
EMBEDDING_MODEL = "text-embedding-3-small"
COLLECTION_NAME = "neustring_deals"

# Global clients & counters
qdrant_client = None
openai_client = None
doc_count    = 0
chunk_count  = 0

# Global UI placeholders
upload_status   = ui.element()
search_results  = ui.element()
stats_container = ui.element()

# Function to get unique document count
def get_unique_document_count():
    unique_filenames = set()
    offset = None
    while True:
        records, next_offset = qdrant_client.scroll(
            collection_name=COLLECTION_NAME,
            limit=100,  # Batch size for efficiency
            with_payload=["filename"],
            with_vectors=False,
            offset=offset
        )
        for record in records:
            filename = record.payload.get("filename")
            if filename:
                unique_filenames.add(filename)
        if next_offset is None:
            break
        offset = next_offset
    return len(unique_filenames)

# Initialize Qdrant and OpenAI clients
def init_clients():
    global qdrant_client, openai_client, doc_count, chunk_count
    if not OPENAI_API_KEY:
        ui.notify("OpenAI API key missing", type="negative")
        return False
    openai_client = OpenAI(api_key=OPENAI_API_KEY)
    qdrant_client = QdrantClient(url=QDRANT_HOST, api_key=QDRANT_API_KEY)
    existing = [c.name for c in qdrant_client.get_collections().collections]
    if COLLECTION_NAME not in existing:
        qdrant_client.create_collection(
            collection_name=COLLECTION_NAME,
            vectors_config=models.VectorParams(
                size=1536, distance=models.Distance.COSINE
            )
        )
        doc_count = 0
        chunk_count = 0
    else:
        # Retrieve counts if collection exists
        count_result = qdrant_client.count(collection_name=COLLECTION_NAME)
        chunk_count = count_result.count
        doc_count = get_unique_document_count()
    return True

# Generate embedding for text using OpenAI
def get_embedding(text: str):
    resp = openai_client.embeddings.create(input=text, model=EMBEDDING_MODEL)
    return resp.data[0].embedding

def extract_text(file_bytes: bytes, mime: str) -> str:
    if mime == "application/pdf":
        reader = PyPDF2.PdfReader(io.BytesIO(file_bytes))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    if mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        doc = docx.Document(io.BytesIO(file_bytes))
        return "\n".join(p.text for p in doc.paragraphs)

    if mime == "text/plain":
        return file_bytes.decode()

    if mime == "application/json":
        try:
            data = json.loads(file_bytes.decode())
            return extract_text_from_json(data)
        except Exception as e:
            return f"[Error reading JSON: {e}]"

    return "[Unsupported file type]"

# Helper to recursively flatten JSON
def extract_text_from_json(data, indent=0) -> str:
    lines = []

    if isinstance(data, dict):
        for key, value in data.items():
            lines.append(" " * indent + f"{key}:")
            lines.append(extract_text_from_json(value, indent + 2))

    elif isinstance(data, list):
        for i, item in enumerate(data):
            lines.append(" " * indent + f"- Item {i}:")
            lines.append(extract_text_from_json(item, indent + 2))

    else:
        lines.append(" " * indent + str(data))

    return "\n".join(lines)

# Process document synchronously (runs in background)
def _process_document(e: events.UploadEventArguments):
    data = e.content.read()
    text = extract_text(data, e.type)
    if not text.strip():
        return None
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=100, length_function=len)
    chunks = splitter.split_text(text)
    points = []
    for i, chunk in enumerate(chunks):
        vec = get_embedding(chunk)
        points.append(PointStruct(
            id=str(uuid.uuid4()),
            vector=vec,
            payload={
                "text": chunk,
                "filename": e.name,
                "chunk_id": i,
                "upload_time": datetime.now().isoformat()
            }
        ))
    qdrant_client.upsert(collection_name=COLLECTION_NAME, points=points)
    return points, len(text), len(points)

def _perform_search(query: str):
    if not qdrant_client or not openai_client:
        if not init_clients():
            return None
    query_embedding = get_embedding(query)
    if not query_embedding:
        return None
    results = qdrant_client.search(
        collection_name=COLLECTION_NAME,
        query_vector=query_embedding,
        limit=5
    )
    context = "\n\n".join([r.payload['text'] for r in results if 'text' in r.payload])
    return results, context

def call_openai_chat(query: str, context: str, doc_names: list) -> str:
    joined_names = ", ".join(doc_names)
    messages = [
        {"role": "system", "content": (
            "You are a helpful assistant. Answer the user's query using the context below. "
            "The context comes from documents: " + joined_names + ". "
            "If any document contains relevant data, use it to answer. "
            "Otherwise, say you don't have enough information."
        )},
        {"role": "user", "content": f"Context:\n{context}\n\nQuestion: {query}"}
    ]
    try:
        response = openai_client.chat.completions.create(
            model="gpt-3.5-turbo",
            messages=messages,
            temperature=0.2,
        )
        return response.choices[0].message.content.strip()
    except Exception as e:
        return f"[OpenAI error: {e}]"

# Async function to handle document uploads
async def upload_document(e: events.UploadEventArguments):
    global doc_count, chunk_count, upload_status
    if qdrant_client is None or openai_client is None:
        if not init_clients():
            return

    # Show processing status immediately
    upload_status.clear()
    with upload_status:
        with ui.card().classes("p-4 bg-blue-50 border-l-4 border-blue-500"):
            with ui.row().classes("items-center gap-3"):
                ui.spinner(size="sm", color="blue")
                ui.label(f"🔄 Processing {e.name}...").classes("text-blue-800 font-medium")

    try:
        # Run heavy processing in a background thread
        loop = asyncio.get_running_loop()
        result = await loop.run_in_executor(None, _process_document, e)
        if result is None:
            upload_status.clear()
            with upload_status:
                with ui.card().classes("p-4 bg-yellow-50 border-l-4 border-yellow-500"):
                    ui.label("⚠️ No text found").classes("text-yellow-800")
            return

        points, text_length, chunk_count_update = result

        # Update counts and UI after processing
        doc_count += 1
        chunk_count += chunk_count_update
        upload_status.clear()
        with upload_status:
            with ui.card().classes("p-4 bg-green-50 border-l-4 border-green-500"):
                with ui.row().classes("items-center gap-3"):
                    ui.icon("check_circle", color="green")
                    with ui.column():
                        ui.label(f"🎉 Processed {e.name}!").classes("text-green-800 font-bold")
                        ui.label(f"📊 {chunk_count_update} chunks added").classes("text-green-600")

        update_stats(doc_count, chunk_count)
    except Exception as e:
        upload_status.clear()
        with upload_status:
            with ui.card().classes('p-6 bg-red-50 border-l-4 border-red-500'):
                ui.label(f'❌ Upload error: {str(e)}').classes('text-red-800')
        ui.notify(f"Upload failed: {str(e)}", type='negative')

# Async function to handle searches
async def search_documents(query: str):
    if not query.strip():
        ui.notify("Please enter a search query", type='warning')
        return

    # Clear previous results and show loading spinner immediately
    search_results.clear()
    with search_results:
        with ui.column().classes('text-center p-8'):
            ui.spinner(size='xl', color='primary')
            ui.label('🔍 Searching through your documents...').classes('text-h6 q-mt-md')

    try:
        # Run search in a background thread
        loop = asyncio.get_running_loop()
        result_pair = await loop.run_in_executor(None, _perform_search, query)
        if not result_pair:
            raise Exception("Failed to perform search")
        results, context = result_pair
        doc_names = [r.payload.get('filename', 'unknown') for r in results]
        llm_answer = await loop.run_in_executor(None, call_openai_chat, query, context, doc_names)

        # Display results after search completes
        search_results.clear()
        with search_results:
            with ui.card().classes('bg-indigo-50 p-4 border-l-4 border-indigo-500'):
                ui.label("🤖 LLM Answer").classes('text-h6 text-indigo-700 font-bold')
                ui.markdown(llm_answer).classes('text-body2 text-indigo-800')
        if not results:
            with search_results:
                with ui.column().classes('text-center p-8 gap-4'):
                    ui.icon('search_off', size='4rem').classes('text-grey-4')
                    ui.label('🤔 No matching documents found').classes('text-h6 text-grey-6')
                    ui.label('Try rephrasing your query or upload more documents').classes('text-grey-5')
            return

        with search_results:
            with ui.row().classes('w-full items-center justify-between q-mb-md'):
                ui.label(f'🎯 Found {len(results)} relevant results').classes('text-h6 font-bold text-primary')
                ui.label(f'Query: "{query}"').classes('text-caption text-grey-6 italic')

            for i, result in enumerate(results):
                with ui.card().classes('result-card p-6 q-mb-md'):
                    with ui.row().classes('w-full items-center justify-between q-mb-sm'):
                        with ui.row().classes('items-center gap-2'):
                            ui.icon('article', color='primary')
                            ui.label(f'Result #{i+1}').classes('text-weight-bold text-primary')
                        score = result.score
                        score_color = 'green' if score > 0.8 else 'orange' if score > 0.6 else 'red'
                        with ui.chip(f'{score:.1%}', icon='trending_up').props(f'color={score_color} text-color=white'):
                            pass

                    with ui.row().classes('items-center gap-4 q-mb-sm'):
                        ui.icon('folder', size='sm', color='grey-6')
                        ui.label(result.payload['filename']).classes('text-weight-medium')
                        ui.separator().props('vertical')
                        ui.icon('schedule', size='sm', color='grey-6')
                        upload_time = datetime.fromisoformat(result.payload['upload_time'])
                        ui.label(upload_time.strftime('%Y-%m-%d %H:%M')).classes('text-grey-6 text-caption')

                    ui.separator().classes('q-my-sm')

                    content = result.payload['text']
                    preview_length = 300

                    if len(content) <= preview_length:
                        ui.markdown(content).classes('text-body2')
                    else:
                        preview_text = content[:preview_length] + "..."
                        ui.markdown(preview_text).classes('text-body2')
                        with ui.expansion('📖 Read full content', icon='expand_more').classes('q-mt-sm'):
                            ui.markdown(content).classes('text-body2 q-pa-md bg-grey-1 rounded')

                    with ui.row().classes('gap-2 q-mt-md'):
                        ui.button('📋 Copy Text',
                                on_click=lambda text=content: ui.run_javascript(f'navigator.clipboard.writeText({repr(text)})')
                        ).props('size=sm color=grey-7 outline')
                        ui.button('🔍 Similar Content',
                                on_click=lambda: search_documents(content[:50] + "...")
                        ).props('size=sm color=secondary outline')

        ui.notify(f"✅ Search completed - {len(results)} results found", type='positive')
    except Exception as e:
        search_results.clear()
        with search_results:
            with ui.card().classes('p-6 bg-red-50 border-l-4 border-red-500'):
                ui.label(f'❌ Search error: {str(e)}').classes('text-red-800')
        ui.notify(f"Search failed: {str(e)}", type='negative')

# Update statistics display
def update_stats(docs: int, chunks: int):
    stats_container.clear()
    with stats_container:
        for icon_name, label, value in [
            ("description", "Documents", docs),
            ("auto_stories", "Chunks", chunks),
            ("sync", "Status", "Ready")
        ]:
            with ui.card().classes("stats-card p-4 text-center min-w-32"):
                ui.icon(icon_name, size="2rem").classes("q-mb-sm")
                ui.label(str(value)).classes("text-h4 font-bold")
                ui.label(label).classes("text-caption")

# Main UI setup
@ui.page("/")
def main():
    global upload_status, search_results, stats_container, doc_count, chunk_count

    # Initialize clients and retrieve counts at startup
    if not init_clients():
        with ui.column().classes('w-full min-h-screen items-center justify-center'):
            ui.label("❌ Failed to initialize clients. Please check your API keys and Qdrant server.").classes('text-h4 text-red-500')
        return

    ui.page_title("Document Vector Search")
    apply_custom_styles()

    with ui.column().classes('w-full min-h-screen'):
        with ui.row().classes('w-full gradient-bg p-8 items-center justify-center'):
            with ui.column().classes('text-center'):
                ui.icon('search', size='3rem').classes('q-mb-md')
                ui.label('🚀 Smart Document Vector Search').classes('text-h3 font-bold q-mb-sm')
                ui.label('Upload • Vectorize • Search • Discover').classes('text-subtitle1 opacity-90 typing-animation')

        with ui.column().classes('w-full max-w-6xl mx-auto p-6 gap-6'):
            stats_container = ui.row().classes('w-full gap-4 q-mb-lg justify-center')
            update_stats(doc_count, chunk_count)

            with ui.column().classes('upload-container'):
                with ui.card().classes("upload-zone p-6 upload-content"):
                    with ui.column().classes("w-full items-center text-center gap-4"):
                        ui.icon('cloud_upload', size='2rem').classes('text-primary')
                        ui.label('📁 Upload Your Documents').classes('text-h5 font-bold text-primary')
                        ui.label('Drag & drop or click to upload PDF, DOCX, or TXT files').classes('text-grey-7')
                    ui.upload(
                        on_upload=upload_document,
                        auto_upload=True,
                        multiple=True
                    ).props('accept=".pdf,.docx,.txt, .json"').classes("w-full q-mt-md justify-center items-center")
                    upload_status = ui.column().classes("w-full q-mt-md")

                with ui.card().classes('search-card p-6'):
                    with ui.row().classes('w-full items-center gap-4 q-mb-md'):
                        ui.icon('manage_search', size='2rem').classes('text-secondary')
                        ui.label('🔍 Search Your Knowledge Base').classes('text-h5 font-bold text-secondary')

                    with ui.row().classes('w-full gap-3'):
                        search_input = ui.input('What would you like to know?').props('outlined dense').classes('flex-grow').style('font-size: 16px;')
                        ui.button('🔎 Search', on_click=lambda: search_documents(search_input.value)).props('color=secondary size=lg push glossy').classes('px-6')

                    with ui.row().classes('gap-2 q-mt-sm'):
                        ui.label('💡 Quick searches:').classes('text-caption text-grey-6')
                        for suggestion in ['summarize', 'key points', 'methodology', 'conclusion']:
                            ui.chip(suggestion, on_click=lambda s=suggestion: setattr(search_input, 'value', s)).props('clickable color=grey-3 text-color=grey-8 size=sm')

                search_results = ui.column().classes('w-full gap-4')

# Run the application
ui.run(title="Document Vector Search", port=8080)